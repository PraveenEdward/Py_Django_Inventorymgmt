import json
from datetime import timedelta

import pandas as pd
from django.db.models import Sum, Q, DecimalField, F, Count
from django.db.models.functions import TruncMonth
from django.urls import reverse_lazy
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from decimal import Decimal

from django.utils import timezone
from openpyxl import Workbook
from django.views.generic import (
    ListView,
    CreateView,
    UpdateView,
    DeleteView
)
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from django.http import JsonResponse, HttpResponse
from .models import (
    Invoice,
    InvoiceItem,
    StockHistory, Settings
)

from .forms import (
    CategoryForm,
    ProductForm, InvoiceForm, SettingsForm,
)

from .forms import ProductExcelUploadForm
from .models import Category, Product



def dashboard(request):

    total_products = Product.objects.count()
    total_categories = Category.objects.count()
    total_invoice = Invoice.objects.count()

    revenue = Invoice.objects.aggregate(
        total=Sum('grand_total')
    )['total'] or 0

    low_stock = Product.objects.filter(stock__lte=5)

    settings = Settings.objects.first()

    # ==========================
    # Monthly Revenue Bar Chart
    # ==========================
    monthly_sales = (
        Invoice.objects
        .annotate(month=TruncMonth('created_at'))
        .values('month')
        .annotate(total=Sum('grand_total'))
        .order_by('month')
    )

    revenue_chart = []

    for sale in monthly_sales:
        revenue_chart.append({
            'y': sale['month'].strftime('%b %Y'),
            'revenue': float(sale['total']),
        })

    # ==========================
    # Top 5 Selling Products Donut Chart
    # ==========================
    top_products = (
        InvoiceItem.objects
        .values('product__product_name')
        .annotate(total_sold=Sum('quantity'))
        .order_by('-total_sold')[:5]
    )

    donut_data = []

    for product in top_products:
        donut_data.append({
            'label': product['product__product_name'],
            'value': product['total_sold']
        })

    # ==========================
    # Context
    # ==========================
    context = {
        'settings': settings,
        'total_products': total_products,
        'total_categories': total_categories,
        'total_invoice': total_invoice,
        'revenue': revenue,
        'low_stock': low_stock,

        # Charts
        'revenue_chart': json.dumps(revenue_chart),
        'sales_donut': json.dumps(donut_data),
    }

    return render(
        request,
        'backend/dashboard.html',
        context
    )


def get_product_by_barcode(request, barcode):

    try:

        product = Product.objects.get(
            barcode=barcode
        )

        return JsonResponse({
            'id': product.id,
            'name': product.product_name,
            'price': float(product.selling_price),
            'stock': product.stock,
        })

    except Product.DoesNotExist:

        return JsonResponse({
            'error': 'Product Not Found'
        })

class CategoryListView(ListView):
    model = Category
    template_name = 'backend/category/list.html'
    context_object_name = 'categories'
    paginate_by = 10

    def get_queryset(self):
        query = self.request.GET.get('q')

        if query:
            return Category.objects.filter(
                category_name__icontains=query
            )

        return Category.objects.all()


class CategoryCreateView(CreateView):
    model = Category
    form_class = CategoryForm
    template_name = 'backend/category/create.html'
    success_url = reverse_lazy('category_list')


class CategoryUpdateView(UpdateView):
    model = Category
    form_class = CategoryForm
    template_name = 'backend/category/update.html'
    success_url = reverse_lazy('category_list')


class CategoryDeleteView(DeleteView):
    model = Category
    template_name = 'backend/category/delete.html'
    success_url = reverse_lazy('category_list')

class ProductListView(ListView):
    model = Product
    template_name = 'backend/product/list.html'
    context_object_name = 'products'
    paginate_by = 10

    def get_queryset(self):

        queryset = Product.objects.select_related(
            'category'
        ).order_by('-id')

        q = self.request.GET.get('q')

        if q:

            queryset = queryset.filter(
                product_name__icontains=q
            )

        return queryset


class ProductCreateView(CreateView):
    model = Product
    form_class = ProductForm
    template_name = 'backend/product/create.html'
    success_url = reverse_lazy('product_list')

    def get_context_data(self, **kwargs):

        context = super().get_context_data(**kwargs)
        context['excel_form'] = ProductExcelUploadForm()

        return context

    def post(self, request, *args, **kwargs):

        if 'excel_upload' in request.POST:

            excel_file = request.FILES.get('excel_file')

            if excel_file:

                df = pd.read_excel(excel_file)

                # Remove extra spaces from Excel headers
                df.columns = df.columns.str.strip()

                for _, row in df.iterrows():
                    category, _ = Category.objects.get_or_create(
                        category_name=str(row['Category']).strip()
                    )

                    Product.objects.update_or_create(
                        barcode=str(row['Barcode']).strip(),
                        defaults={
                            'category': category,
                            'product_name': str(row['Product Name']).strip(),
                            'purchase_price': row['Purchase Price'],
                            'selling_price': row['Selling Price'],
                            'stock': int(row['Stock']),
                            'minimum_stock': int(
                                row.get('Minimum Stock', 5)
                            ),
                        }
                    )

                messages.success(
                    request,
                    "Products Imported Successfully"
                )

                return redirect('product_list')

        return super().post(request, *args, **kwargs)


class ProductUpdateView(UpdateView):
    model = Product
    form_class = ProductForm
    template_name = 'backend/product/update.html'
    success_url = reverse_lazy('product_list')


class ProductDeleteView(DeleteView):
    model = Product
    template_name = 'backend/product/delete.html'
    success_url = reverse_lazy('product_list')

class InvoiceListView(ListView):
    model = Invoice
    template_name = 'backend/invoice/list.html'
    context_object_name = 'invoices'
    ordering = ['-id']
    paginate_by = 10

    def get_queryset(self):
        queryset = Invoice.objects.order_by('-id')

        q = self.request.GET.get('q')

        if q:
            queryset = queryset.filter(
                Q(invoice_no__icontains=q) |
                Q(customer_name__icontains=q) |
                Q(customer_mobile__icontains=q)
            )

        return queryset


def create_invoice(request):

    products = Product.objects.all()

    if request.method == 'POST':

        form = InvoiceForm(request.POST)

        if form.is_valid():

            invoice = form.save(commit=False)

            total_amount = 0

            invoice.save()

            product_ids = request.POST.getlist('product[]')
            quantities = request.POST.getlist('qty[]')

            for product_id, qty in zip(product_ids, quantities):

                if not product_id:
                    continue

                product = Product.objects.get(
                    id=product_id
                )

                qty = int(qty)

                # Check Stock
                if product.stock < qty:

                    messages.error(
                        request,
                        f'{product.product_name} stock not available'
                    )

                    invoice.delete()

                    return redirect('create_invoice')

                # Create Invoice Item
                InvoiceItem.objects.create(
                    invoice=invoice,
                    product=product,
                    quantity=qty,
                    price=product.selling_price,
                    subtotal=product.selling_price * qty
                )

                # Reduce Stock
                product.stock -= qty
                product.save()

                total_amount += (
                    product.selling_price * qty
                )

            discount = invoice.discount_amount or 0

            invoice.total_amount = total_amount
            invoice.grand_total = total_amount - discount

            invoice.save()

            messages.success(
                request,
                'Invoice Created Successfully'
            )

            return redirect('invoice_list')

    else:

        form = InvoiceForm()

    return render(
        request,
        'backend/invoice/create.html',
        {
            'form': form,
            'products': products,
        }
    )

def update_invoice(request, pk):

    invoice = get_object_or_404(
        Invoice,
        pk=pk
    )

    if request.method == 'POST':

        # Restore stock
        for item in invoice.items.all():

            product = item.product

            product.stock += item.quantity

            product.save()

        invoice.items.all().delete()

        invoice.customer_name = request.POST.get(
            'customer_name'
        )

        invoice.customer_mobile = request.POST.get(
            'customer_mobile'
        )

        invoice.payment_method = request.POST.get(
            'payment_method'
        )

        invoice.discount_amount = Decimal(
            request.POST.get(
                'discount_amount',
                0
            ) or 0
        )

        invoice.save()

        product_ids = request.POST.getlist(
            'product[]'
        )

        quantities = request.POST.getlist(
            'qty[]'
        )

        for product_id, qty in zip(
            product_ids,
            quantities
        ):

            if product_id:

                product = Product.objects.get(
                    pk=product_id
                )

                InvoiceItem.objects.create(
                    invoice=invoice,
                    product=product,
                    quantity=int(qty)
                )

        invoice.total_qty = (
            invoice.items.aggregate(
                qty=Sum('quantity')
            )['qty'] or 0
        )

        invoice.total_amount = (
            invoice.items.aggregate(
                total=Sum('subtotal')
            )['total'] or Decimal('0.00')
        )

        invoice.grand_total = (
            invoice.total_amount -
            invoice.discount_amount
        )

        invoice.save()

        messages.success(
            request,
            'Invoice Updated Successfully'
        )

        return redirect(
            'invoice_list'
        )

    products = Product.objects.all()

    return render(
        request,
        'backend/invoice/update.html',
        {
            'invoice': invoice,
            'products': products,
        }
    )

def delete_invoice(request, pk):

    invoice = get_object_or_404(
        Invoice,
        pk=pk
    )

    if request.method == 'POST':

        for item in invoice.items.all():

            product = item.product

            product.stock += item.quantity

            product.save()

        invoice.delete()

        messages.success(
            request,
            'Invoice Deleted Successfully'
        )

        return redirect(
            'invoice_list'
        )

    return render(
        request,
        'backend/invoice/delete.html',
        {
            'invoice': invoice
        }
    )

class StockHistoryListView(ListView):
    model = StockHistory
    template_name = 'backend/stock/list.html'
    context_object_name = 'stocks'
    paginate_by = 10

    def get_queryset(self):

        queryset = StockHistory.objects.select_related(
            'product',
            'product__category'
        ).order_by('-created_at')

        q = self.request.GET.get('q')

        category = self.request.GET.get('category')

        if q:
            queryset = queryset.filter(
                Q(product__product_name__icontains=q) |
                Q(product__barcode__icontains=q)
            )

        if category:
            queryset = queryset.filter(
                product__category_id=category
            )

        return queryset

    def get_context_data(self, **kwargs):

        context = super().get_context_data(**kwargs)

        context['categories'] = Category.objects.all()

        return context

def invoice_pdf(request, pk):

    invoice = Invoice.objects.get(pk=pk)

    response = HttpResponse(
        content_type='application/pdf'
    )

    response['Content-Disposition'] = (
        f'attachment; filename="{invoice.invoice_no}.pdf"'
    )

    pdf = canvas.Canvas(
        response,
        pagesize=A4
    )

    width, height = A4

    # Shop Details
    pdf.setFont("Helvetica-Bold", 18)
    pdf.drawString(50, 800, "POS BILLING SYSTEM")

    pdf.setFont("Helvetica", 12)
    pdf.drawString(
        50,
        770,
        f"Invoice No : {invoice.invoice_no}"
    )

    pdf.drawString(
        50,
        750,
        f"Customer : {invoice.customer_name}"
    )

    pdf.drawString(
        50,
        730,
        f"Mobile : {invoice.customer_mobile}"
    )

    pdf.drawString(
        50,
        710,
        f"Payment : {invoice.payment_method}"
    )

    # Table Header
    y = 650

    pdf.setFont("Helvetica-Bold", 12)

    pdf.drawString(50, y, "Product")
    pdf.drawString(250, y, "Price")
    pdf.drawString(330, y, "Qty")
    pdf.drawString(420, y, "Subtotal")

    y -= 20

    pdf.line(50, y, 550, y)

    y -= 20

    # Invoice Items
    pdf.setFont("Helvetica", 11)

    for item in invoice.items.all():

        pdf.drawString(
            50,
            y,
            item.product.product_name
        )

        pdf.drawString(
            250,
            y,
            str(item.price)
        )

        pdf.drawString(
            330,
            y,
            str(item.quantity)
        )

        pdf.drawString(
            420,
            y,
            str(item.subtotal)
        )

        y -= 20

    y -= 20

    pdf.line(50, y, 550, y)

    y -= 30

    pdf.setFont("Helvetica-Bold", 12)

    pdf.drawString(
        350,
        y,
        f"Total Qty : {invoice.total_qty}"
    )

    y -= 25

    pdf.drawString(
        350,
        y,
        f"Total : ₹{invoice.total_amount}"
    )

    y -= 25

    pdf.drawString(
        350,
        y,
        f"Discount : ₹{invoice.discount_amount}"
    )

    y -= 25

    pdf.drawString(
        350,
        y,
        f"Grand Total : ₹{invoice.grand_total}"
    )

    pdf.save()

    return response

class SettingsListView(ListView):

    model = Settings
    template_name = 'backend/settings/list.html'
    context_object_name = 'settings_list'
    paginate_by = 10
    ordering = ['-id']

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)

        context['setting'] = Settings.objects.first()

        return context

class SettingsCreateView(CreateView):

    model = Settings
    form_class = SettingsForm
    template_name = 'backend/settings/create.html'
    success_url = reverse_lazy('settings_list')

class SettingsUpdateView(UpdateView):

    model = Settings
    form_class = SettingsForm
    template_name = 'backend/settings/update.html'
    success_url = reverse_lazy('settings_list')

class SettingsDeleteView(DeleteView):

    model = Settings
    template_name = 'backend/settings/delete.html'
    success_url = reverse_lazy('settings_list')

def sales_report(request):

    invoices = Invoice.objects.all().order_by('-created_at')

    from_date = request.GET.get('from_date')
    to_date = request.GET.get('to_date')
    export = request.GET.get('export')

    if from_date and to_date:
        invoices = invoices.filter(
            created_at__date__range=[from_date, to_date]
        )

    total_sales = invoices.aggregate(
        total=Sum('grand_total')
    )['total'] or 0

    # Excel Export
    if export == 'excel':

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )

        response['Content-Disposition'] = (
            'attachment; filename="sales_report.xlsx"'
        )

        wb = Workbook()
        ws = wb.active
        ws.title = 'Sales Report'

        ws.append([
            'Invoice No',
            'Date',
            'Customer',
            'Mobile',
            'Payment',
            'Grand Total'
        ])

        for invoice in invoices:
            ws.append([
                invoice.invoice_no,
                invoice.created_at.strftime('%d-%m-%Y'),
                invoice.customer_name,
                invoice.customer_mobile,
                invoice.payment_method,
                float(invoice.grand_total),
            ])

        ws.append([])
        ws.append([
            '',
            '',
            '',
            '',
            'Total Sales',
            float(total_sales)
        ])

        wb.save(response)

        return response

    # PDF Export
    if export == 'pdf':

        response = HttpResponse(
            content_type='application/pdf'
        )

        response['Content-Disposition'] = (
            'attachment; filename="sales_report.pdf"'
        )

        p = canvas.Canvas(response)

        p.setFont("Helvetica-Bold", 14)
        p.drawString(200, 800, "Sales Report")

        y = 760

        p.setFont("Helvetica", 10)

        for invoice in invoices:

            p.drawString(
                40,
                y,
                f"{invoice.invoice_no}"
            )

            p.drawString(
                140,
                y,
                invoice.created_at.strftime('%d-%m-%Y')
            )

            p.drawString(
                230,
                y,
                invoice.customer_name or 'Walk-in'
            )

            p.drawString(
                340,
                y,
                invoice.payment_method
            )

            p.drawString(
                430,
                y,
                f"₹{invoice.grand_total}"
            )

            y -= 20

            if y < 50:
                p.showPage()
                y = 800

        p.drawString(
            330,
            y - 30,
            f"Total Sales : ₹{total_sales}"
        )

        p.save()

        return response

    return render(
        request,
        'backend/reports/sales_report.html',
        {
            'invoices': invoices,
            'total_sales': total_sales,
            'from_date': from_date,
            'to_date': to_date,
        }
    )

def product_sales_report(request):

    items = InvoiceItem.objects.all()

    from_date = request.GET.get('from_date')
    to_date = request.GET.get('to_date')
    export = request.GET.get('export')

    if from_date and to_date:
        items = items.filter(
            invoice__created_at__date__range=[
                from_date,
                to_date
            ]
        )

    report = items.values(
        'product__product_name'
    ).annotate(
        total_qty=Sum('quantity'),
        total_amount=Sum('subtotal')
    )

    # Excel Export
    if export == 'excel':

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )

        response[
            'Content-Disposition'
        ] = 'attachment; filename=product_sales_report.xlsx'

        wb = Workbook()
        ws = wb.active
        ws.title = 'Product Sales'

        ws.append([
            'S.No',
            'Product Name',
            'Quantity Sold',
            'Total Amount'
        ])

        count = 1

        for item in report:

            ws.append([
                count,
                item['product__product_name'],
                item['total_qty'],
                float(item['total_amount']),
            ])

            count += 1

        wb.save(response)

        return response

    # PDF Export
    if export == 'pdf':

        response = HttpResponse(
            content_type='application/pdf'
        )

        response[
            'Content-Disposition'
        ] = 'attachment; filename=product_sales_report.pdf'

        p = canvas.Canvas(response)

        p.setFont(
            "Helvetica-Bold",
            14
        )

        p.drawString(
            180,
            800,
            "Product Sales Report"
        )

        y = 760

        p.setFont(
            "Helvetica-Bold",
            10
        )

        p.drawString(40, y, "S.No")
        p.drawString(90, y, "Product")
        p.drawString(300, y, "Qty")
        p.drawString(380, y, "Amount")

        y -= 20

        p.setFont(
            "Helvetica",
            10
        )

        count = 1

        for item in report:

            p.drawString(
                40,
                y,
                str(count)
            )

            p.drawString(
                90,
                y,
                str(item['product__product_name'])
            )

            p.drawString(
                300,
                y,
                str(item['total_qty'])
            )

            p.drawString(
                380,
                y,
                "₹{}".format(item['total_amount'])
            )

            y -= 20

            count += 1

            if y <= 50:
                p.showPage()
                y = 800

        p.save()

        return response

    return render(
        request,
        'backend/reports/product_sales_report.html',
        {
            'report': report,
            'from_date': from_date,
            'to_date': to_date,
        }
    )


def stock_report(request):

    histories = StockHistory.objects.select_related(
        'product'
    ).order_by('-created_at')

    from_date = request.GET.get('from_date', '')
    to_date = request.GET.get('to_date', '')
    export = request.GET.get('export')

    if from_date and to_date:
        histories = histories.filter(
            created_at__date__range=[
                from_date,
                to_date
            ]
        )

    total_products = histories.values(
        'product'
    ).distinct().count()

    total_stock_in = histories.filter(
        action='IN'
    ).aggregate(
        total=Sum('quantity')
    )['total'] or 0

    total_stock_out = histories.filter(
        action='OUT'
    ).aggregate(
        total=Sum('quantity')
    )['total'] or 0

    # =====================
    # Excel Export
    # =====================
    if export == 'excel':

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )

        response[
            'Content-Disposition'
        ] = 'attachment; filename="stock_report.xlsx"'

        workbook = Workbook()

        worksheet = workbook.active

        worksheet.title = 'Stock Report'

        worksheet.append([
            'S.No',
            'Date',
            'Product Name',
            'Previous Stock',
            'Added/Removed',
            'Current Stock',
            'Action',
            'Remarks',
        ])

        for index, history in enumerate(histories, start=1):

            worksheet.append([
                index,
                history.created_at.strftime('%d-%m-%Y %H:%M'),
                history.product.product_name,
                history.previous_stock,
                history.quantity,
                history.current_stock,
                history.action,
                history.remarks,
            ])

        worksheet.append([])

        worksheet.append([
            '',
            '',
            'Total Products',
            total_products,
            '',
            '',
            '',
            '',
        ])

        workbook.save(response)

        return response

    # =====================
    # PDF Export
    # =====================
    if export == 'pdf':

        response = HttpResponse(
            content_type='application/pdf'
        )

        response[
            'Content-Disposition'
        ] = 'attachment; filename="stock_report.pdf"'

        pdf = canvas.Canvas(response)

        pdf.setTitle('Stock Report')

        pdf.setFont(
            'Helvetica-Bold',
            16
        )

        pdf.drawString(
            220,
            800,
            'Stock Report'
        )

        y = 760

        pdf.setFont(
            'Helvetica-Bold',
            8
        )

        pdf.drawString(20, y, 'No')
        pdf.drawString(45, y, 'Date')
        pdf.drawString(120, y, 'Product')
        pdf.drawString(240, y, 'Prev')
        pdf.drawString(290, y, 'Qty')
        pdf.drawString(340, y, 'Current')
        pdf.drawString(410, y, 'Action')
        pdf.drawString(470, y, 'Remarks')

        y -= 20

        pdf.setFont(
            'Helvetica',
            8
        )

        for index, history in enumerate(histories, start=1):

            pdf.drawString(
                20,
                y,
                str(index)
            )

            pdf.drawString(
                45,
                y,
                history.created_at.strftime('%d-%m-%Y')
            )

            pdf.drawString(
                120,
                y,
                history.product.product_name[:18]
            )

            pdf.drawString(
                240,
                y,
                str(history.previous_stock)
            )

            pdf.drawString(
                290,
                y,
                str(history.quantity)
            )

            pdf.drawString(
                340,
                y,
                str(history.current_stock)
            )

            pdf.drawString(
                410,
                y,
                history.action
            )

            pdf.drawString(
                470,
                y,
                history.remarks[:15]
            )

            y -= 18

            if y <= 40:

                pdf.showPage()

                y = 760

                pdf.setFont(
                    'Helvetica-Bold',
                    8
                )

                pdf.drawString(20, y, 'No')
                pdf.drawString(45, y, 'Date')
                pdf.drawString(120, y, 'Product')
                pdf.drawString(240, y, 'Prev')
                pdf.drawString(290, y, 'Qty')
                pdf.drawString(340, y, 'Current')
                pdf.drawString(410, y, 'Action')
                pdf.drawString(470, y, 'Remarks')

                y -= 20

                pdf.setFont(
                    'Helvetica',
                    8
                )

        pdf.setFont(
            'Helvetica-Bold',
            10
        )

        pdf.drawString(
            20,
            y - 20,
            f'Total Products : {total_products}'
        )

        pdf.drawString(
            200,
            y - 20,
            f'Stock IN : {total_stock_in}'
        )

        pdf.drawString(
            350,
            y - 20,
            f'Stock OUT : {total_stock_out}'
        )

        pdf.save()

        return response

    return render(
        request,
        'backend/reports/stock_report.html',
        {
            'histories': histories,
            'from_date': from_date,
            'to_date': to_date,
            'total_products': total_products,
            'total_stock_in': total_stock_in,
            'total_stock_out': total_stock_out,
        }
    )

def profit_report(request):

    items = InvoiceItem.objects.select_related(
        'invoice',
        'product'
    ).all()

    from_date = request.GET.get('from_date', '')
    to_date = request.GET.get('to_date', '')
    export = request.GET.get('export')

    # Date Filter
    if (
        from_date and
        to_date and
        from_date != 'None' and
        to_date != 'None'
    ):
        items = items.filter(
            invoice__created_at__date__range=(
                from_date,
                to_date
            )
        )

    # Total Profit
    total_profit = items.aggregate(
        profit=Sum(
            (
                F('price') -
                F('product__purchase_price')
            ) * F('quantity'),
            output_field=DecimalField(
                max_digits=15,
                decimal_places=2
            )
        )
    )['profit'] or 0

    # ======================
    # Excel Export
    # ======================
    if export == 'excel':

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )

        response[
            'Content-Disposition'
        ] = 'attachment; filename="profit_report.xlsx"'

        workbook = Workbook()

        worksheet = workbook.active
        worksheet.title = 'Profit Report'

        worksheet.append([
            'S.No',
            'Invoice No',
            'Date',
            'Product',
            'Quantity',
            'Purchase Price',
            'Selling Price',
            'Profit',
        ])

        for index, item in enumerate(items, start=1):

            profit = (
                item.price -
                item.product.purchase_price
            ) * item.quantity

            worksheet.append([
                index,
                item.invoice.invoice_no,
                item.invoice.created_at.strftime('%d-%m-%Y'),
                item.product.product_name,
                item.quantity,
                float(item.product.purchase_price),
                float(item.price),
                float(profit),
            ])

        worksheet.append([])

        worksheet.append([
            '',
            '',
            '',
            '',
            '',
            '',
            'Total Profit',
            float(total_profit),
        ])

        workbook.save(response)

        return response

    # ======================
    # PDF Export
    # ======================
    if export == 'pdf':

        response = HttpResponse(
            content_type='application/pdf'
        )

        response[
            'Content-Disposition'
        ] = 'attachment; filename="profit_report.pdf"'

        pdf = canvas.Canvas(response)

        pdf.setTitle('Profit Report')

        pdf.setFont(
            'Helvetica-Bold',
            16
        )

        pdf.drawString(
            220,
            800,
            'Profit Report'
        )

        y = 760

        pdf.setFont(
            'Helvetica-Bold',
            8
        )

        pdf.drawString(20, y, 'No')
        pdf.drawString(45, y, 'Invoice')
        pdf.drawString(120, y, 'Date')
        pdf.drawString(190, y, 'Product')
        pdf.drawString(300, y, 'Qty')
        pdf.drawString(340, y, 'Buy')
        pdf.drawString(390, y, 'Sell')
        pdf.drawString(450, y, 'Profit')

        y -= 20

        pdf.setFont(
            'Helvetica',
            8
        )

        for index, item in enumerate(items, start=1):

            profit = (
                item.price -
                item.product.purchase_price
            ) * item.quantity

            pdf.drawString(
                20,
                y,
                str(index)
            )

            pdf.drawString(
                45,
                y,
                item.invoice.invoice_no
            )

            pdf.drawString(
                120,
                y,
                item.invoice.created_at.strftime(
                    '%d-%m-%Y'
                )
            )

            pdf.drawString(
                190,
                y,
                item.product.product_name[:20]
            )

            pdf.drawString(
                300,
                y,
                str(item.quantity)
            )

            pdf.drawString(
                340,
                y,
                str(item.product.purchase_price)
            )

            pdf.drawString(
                390,
                y,
                str(item.price)
            )

            pdf.drawString(
                450,
                y,
                str(profit)
            )

            y -= 18

            if y <= 40:

                pdf.showPage()

                y = 760

                pdf.setFont(
                    'Helvetica-Bold',
                    8
                )

                pdf.drawString(20, y, 'No')
                pdf.drawString(45, y, 'Invoice')
                pdf.drawString(120, y, 'Date')
                pdf.drawString(190, y, 'Product')
                pdf.drawString(300, y, 'Qty')
                pdf.drawString(340, y, 'Buy')
                pdf.drawString(390, y, 'Sell')
                pdf.drawString(450, y, 'Profit')

                y -= 20

                pdf.setFont(
                    'Helvetica',
                    8
                )

        pdf.setFont(
            'Helvetica-Bold',
            10
        )

        pdf.drawString(
            320,
            y - 20,
            f"Total Profit : ₹{total_profit}"
        )

        pdf.save()

        return response

    # Profit for HTML Table
    for item in items:
        item.profit = (
            item.price -
            item.product.purchase_price
        ) * item.quantity

    return render(
        request,
        'backend/reports/profit_report.html',
        {
            'items': items,
            'total_profit': total_profit,
            'from_date': from_date,
            'to_date': to_date,
        }
    )


def customer_report(request):

    invoices = Invoice.objects.all().order_by(
        '-created_at'
    )

    from_date = request.GET.get(
        'from_date',
        ''
    )

    to_date = request.GET.get(
        'to_date',
        ''
    )

    export = request.GET.get(
        'export'
    )

    # Date Filter
    if (
        from_date and
        to_date and
        from_date != 'None' and
        to_date != 'None'
    ):

        invoices = invoices.filter(
            created_at__date__range=(
                from_date,
                to_date
            )
        )

    total_sales = invoices.aggregate(
        total=Sum('grand_total')
    )['total'] or 0

    total_customers = invoices.count()

    # ==========================
    # Excel Export
    # ==========================
    if export == 'excel':

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )

        response[
            'Content-Disposition'
        ] = 'attachment; filename="customer_report.xlsx"'

        workbook = Workbook()

        worksheet = workbook.active
        worksheet.title = 'Customer Report'

        worksheet.append([
            'S.No',
            'Invoice No',
            'Date',
            'Customer Name',
            'Mobile',
            'Payment Method',
            'Grand Total',
        ])

        for index, invoice in enumerate(
            invoices,
            start=1
        ):

            worksheet.append([
                index,
                invoice.invoice_no,
                invoice.created_at.strftime(
                    '%d-%m-%Y'
                ),
                invoice.customer_name,
                invoice.customer_mobile,
                invoice.payment_method,
                float(invoice.grand_total),
            ])

        worksheet.append([])

        worksheet.append([
            '',
            '',
            '',
            '',
            '',
            'Total Sales',
            float(total_sales),
        ])

        workbook.save(response)

        return response

    # ==========================
    # PDF Export
    # ==========================
    if export == 'pdf':

        response = HttpResponse(
            content_type='application/pdf'
        )

        response[
            'Content-Disposition'
        ] = 'attachment; filename="customer_report.pdf"'

        pdf = canvas.Canvas(response)

        pdf.setTitle(
            'Customer Report'
        )

        pdf.setFont(
            'Helvetica-Bold',
            16
        )

        pdf.drawString(
            220,
            800,
            'Customer Report'
        )

        y = 760

        pdf.setFont(
            'Helvetica-Bold',
            8
        )

        pdf.drawString(20, y, 'No')
        pdf.drawString(45, y, 'Invoice')
        pdf.drawString(120, y, 'Date')
        pdf.drawString(180, y, 'Customer')
        pdf.drawString(280, y, 'Mobile')
        pdf.drawString(360, y, 'Payment')
        pdf.drawString(450, y, 'Amount')

        y -= 20

        pdf.setFont(
            'Helvetica',
            8
        )

        for index, invoice in enumerate(
            invoices,
            start=1
        ):

            pdf.drawString(
                20,
                y,
                str(index)
            )

            pdf.drawString(
                45,
                y,
                invoice.invoice_no
            )

            pdf.drawString(
                120,
                y,
                invoice.created_at.strftime(
                    '%d-%m-%Y'
                )
            )

            pdf.drawString(
                180,
                y,
                invoice.customer_name[:15]
            )

            pdf.drawString(
                280,
                y,
                invoice.customer_mobile
            )

            pdf.drawString(
                360,
                y,
                invoice.payment_method
            )

            pdf.drawString(
                450,
                y,
                str(invoice.grand_total)
            )

            y -= 18

            if y <= 40:

                pdf.showPage()

                y = 760

                pdf.setFont(
                    'Helvetica-Bold',
                    8
                )

                pdf.drawString(20, y, 'No')
                pdf.drawString(45, y, 'Invoice')
                pdf.drawString(120, y, 'Date')
                pdf.drawString(180, y, 'Customer')
                pdf.drawString(280, y, 'Mobile')
                pdf.drawString(360, y, 'Payment')
                pdf.drawString(450, y, 'Amount')

                y -= 20

                pdf.setFont(
                    'Helvetica',
                    8
                )

        pdf.setFont(
            'Helvetica-Bold',
            10
        )

        pdf.drawString(
            320,
            y - 20,
            f"Total Sales : ₹{total_sales}"
        )

        pdf.save()

        return response

    return render(
        request,
        'backend/reports/customer_report.html',
        {
            'invoices': invoices,
            'from_date': from_date,
            'to_date': to_date,
            'total_sales': total_sales,
            'total_customers': total_customers,
        }
    )

def revenue_report(request):

    today = timezone.now().date()

    invoices = Invoice.objects.all().order_by('-created_at')

    from_date = request.GET.get('from_date', '')
    to_date = request.GET.get('to_date', '')
    export = request.GET.get('export')

    if (
        from_date and
        to_date and
        from_date != 'None' and
        to_date != 'None'
    ):
        invoices = invoices.filter(
            created_at__date__range=(
                from_date,
                to_date
            )
        )

    weekly_revenue = Invoice.objects.filter(
        created_at__date__gte=today - timedelta(days=7)
    ).aggregate(
        total=Sum('grand_total')
    )['total'] or 0

    monthly_revenue = Invoice.objects.filter(
        created_at__year=today.year,
        created_at__month=today.month
    ).aggregate(
        total=Sum('grand_total')
    )['total'] or 0

    yearly_revenue = Invoice.objects.filter(
        created_at__year=today.year
    ).aggregate(
        total=Sum('grand_total')
    )['total'] or 0

    total_revenue = invoices.aggregate(
        total=Sum('grand_total')
    )['total'] or 0

    if export == 'excel':

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )

        response['Content-Disposition'] = (
            'attachment; filename="revenue_report.xlsx"'
        )

        workbook = Workbook()

        worksheet = workbook.active
        worksheet.title = 'Revenue Report'

        worksheet.append([
            'S.No',
            'Invoice No',
            'Date',
            'Customer',
            'Revenue',
        ])

        for index, invoice in enumerate(
            invoices,
            start=1
        ):
            worksheet.append([
                index,
                invoice.invoice_no,
                invoice.created_at.strftime('%d-%m-%Y'),
                invoice.customer_name,
                float(invoice.grand_total),
            ])

        worksheet.append([])

        worksheet.append([
            '',
            '',
            '',
            'Total Revenue',
            float(total_revenue),
        ])

        workbook.save(response)

        return response

    if export == 'excel':

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )

        response['Content-Disposition'] = (
            'attachment; filename="revenue_report.xlsx"'
        )

        workbook = Workbook()

        worksheet = workbook.active
        worksheet.title = 'Revenue Report'

        worksheet.append([
            'S.No',
            'Invoice No',
            'Date',
            'Customer',
            'Revenue',
        ])

        for index, invoice in enumerate(
            invoices,
            start=1
        ):
            worksheet.append([
                index,
                invoice.invoice_no,
                invoice.created_at.strftime('%d-%m-%Y'),
                invoice.customer_name,
                float(invoice.grand_total),
            ])

        worksheet.append([])

        worksheet.append([
            '',
            '',
            '',
            'Total Revenue',
            float(total_revenue),
        ])

        workbook.save(response)

        return response

    return render(
        request,
        'backend/reports/revenue_report.html',
        {
            'invoices': invoices,
            'weekly_revenue': weekly_revenue,
            'monthly_revenue': monthly_revenue,
            'yearly_revenue': yearly_revenue,
            'total_revenue': total_revenue,
            'from_date': from_date,
            'to_date': to_date,
        }
    )
